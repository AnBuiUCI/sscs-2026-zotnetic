import sys
from pathlib import Path
sys.path.insert(0, "/foss/designs/a_zonetic2026/reportes")
from pptx import Presentation
def tits(f):
    out=[]
    for i,s in enumerate(Presentation(f).slides,1):
        t=""
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                t=sh.text_frame.text.strip().splitlines()[0][:70]; break
        out.append((i,t))
    return out
for i,t in tits(sys.argv[1]): print(f"{i:3d}  {t}")
